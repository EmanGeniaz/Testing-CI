"""PPTX report builder — generates polished PowerPoint presentations from
tagged analysis data and structured report JSON.

Colour palette
--------------
- Primary purple:  #6C4CFF
- Dark navy text:  #14132A
- Light background: #FAFAFF
- Accent grey:     #E8E6F0
- White:           #FFFFFF
"""

import io
import re
from datetime import datetime

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# ── Colour constants ──────────────────────────────────────────────────────────

CLR_PRIMARY    = RGBColor(0x6C, 0x4C, 0xFF)   # purple
CLR_DARK       = RGBColor(0x14, 0x13, 0x2A)   # navy
CLR_LIGHT_BG   = RGBColor(0xFA, 0xFA, 0xFF)   # near-white
CLR_ACCENT     = RGBColor(0xE8, 0xE6, 0xF0)   # light grey
CLR_WHITE      = RGBColor(0xFF, 0xFF, 0xFF)
CLR_GREEN      = RGBColor(0x27, 0xAE, 0x60)
CLR_ORANGE     = RGBColor(0xF3, 0x9C, 0x12)
CLR_RED        = RGBColor(0xE7, 0x4C, 0x3C)
CLR_MID_GREY   = RGBColor(0x66, 0x66, 0x66)

FONT_FAMILY = "Calibri"

SLIDE_WIDTH  = Inches(13.333)   # 16:9 widescreen
SLIDE_HEIGHT = Inches(7.5)


# ── Helpers ───────────────────────────────────────────────────────────────────

def _strip_md(text: str) -> str:
    """Remove markdown bold / italic markers for plain-text contexts."""
    text = re.sub(r"\*\*(.+?)\*\*", r"\1", text)
    text = re.sub(r"\*(.+?)\*", r"\1", text)
    return text


def _set_font(run, size: int = 12, bold: bool = False, italic: bool = False,
              color: RGBColor = CLR_DARK, name: str = FONT_FAMILY):
    run.font.size  = Pt(size)
    run.font.bold  = bold
    run.font.italic = italic
    run.font.color.rgb = color
    run.font.name  = name


def _add_textbox(slide, left, top, width, height, text: str,
                 font_size: int = 12, bold: bool = False, color: RGBColor = CLR_DARK,
                 alignment=PP_ALIGN.LEFT, word_wrap: bool = True):
    """Utility — add a single-paragraph text box and return the shape."""
    txbox = slide.shapes.add_textbox(left, top, width, height)
    txbox.word_wrap = word_wrap
    tf = txbox.text_frame
    tf.word_wrap = word_wrap
    p = tf.paragraphs[0]
    p.alignment = alignment
    run = p.add_run()
    run.text = text
    _set_font(run, size=font_size, bold=bold, color=color)
    return txbox


def _add_rich_body(text_frame, text: str, font_size: int = 12,
                   color: RGBColor = CLR_DARK):
    """Parse **bold** and *italic* markdown in *text* and add runs to the
    first paragraph of *text_frame*."""
    p = text_frame.paragraphs[0]
    p.space_after = Pt(4)
    # Split by bold/italic markers
    tokens = re.split(r"(\*\*.*?\*\*|\*.*?\*)", text)
    for token in tokens:
        if token.startswith("**") and token.endswith("**"):
            run = p.add_run()
            run.text = token[2:-2]
            _set_font(run, size=font_size, bold=True, color=color)
        elif token.startswith("*") and token.endswith("*"):
            run = p.add_run()
            run.text = token[1:-1]
            _set_font(run, size=font_size, italic=True, color=color)
        else:
            run = p.add_run()
            run.text = token
            _set_font(run, size=font_size, color=color)


def _confidence_color(conf: str) -> RGBColor:
    return {"high": CLR_GREEN, "medium": CLR_ORANGE, "low": CLR_RED}.get(conf, CLR_MID_GREY)


def _set_slide_bg(slide, color: RGBColor = CLR_LIGHT_BG):
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = color


def _add_purple_bar(slide):
    """Add a thin purple accent bar at the top of the slide."""
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0), Inches(0),
        SLIDE_WIDTH, Inches(0.12),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = CLR_PRIMARY
    shape.line.fill.background()


# ── Slide builders ────────────────────────────────────────────────────────────

def _build_title_slide(prs, report: dict):
    """Slide 1 — Title slide with report title, subtitle, and metadata."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank layout
    _set_slide_bg(slide, CLR_DARK)

    # Purple block at top
    block = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0), Inches(0),
        SLIDE_WIDTH, Inches(3.0),
    )
    block.fill.solid()
    block.fill.fore_color.rgb = CLR_PRIMARY
    block.line.fill.background()

    # Title
    _add_textbox(slide, Inches(0.8), Inches(0.6), Inches(11.5), Inches(1.5),
                 report.get("title", "Analysis Report"),
                 font_size=36, bold=True, color=CLR_WHITE,
                 alignment=PP_ALIGN.LEFT)

    # Subtitle
    _add_textbox(slide, Inches(0.8), Inches(2.1), Inches(11.5), Inches(0.7),
                 report.get("subtitle", ""),
                 font_size=18, color=CLR_WHITE, alignment=PP_ALIGN.LEFT)

    # Metadata block
    meta = report.get("metadata", {})
    meta_lines = []
    if meta.get("generated_at"):
        try:
            dt = datetime.fromisoformat(meta["generated_at"])
            meta_lines.append(f"Generated: {dt.strftime('%d %B %Y')}")
        except (ValueError, TypeError):
            meta_lines.append(f"Generated: {meta['generated_at']}")
    if meta.get("items_reviewed"):
        meta_lines.append(f"Items reviewed: {meta['items_reviewed']}")
    if meta.get("report_type"):
        meta_lines.append(f"Report type: {meta['report_type']}")
    if meta.get("method"):
        meta_lines.append(f"Method: {meta['method']}")

    meta_text = "   |   ".join(meta_lines) if meta_lines else ""
    _add_textbox(slide, Inches(0.8), Inches(3.5), Inches(11.5), Inches(0.5),
                 meta_text, font_size=13, color=CLR_ACCENT,
                 alignment=PP_ALIGN.LEFT)


def _build_exec_summary_slide(prs, report: dict):
    """Slide 2 — Executive Summary (first section body)."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _set_slide_bg(slide)
    _add_purple_bar(slide)

    _add_textbox(slide, Inches(0.8), Inches(0.4), Inches(11.5), Inches(0.7),
                 "Executive Summary", font_size=28, bold=True, color=CLR_PRIMARY)

    sections = report.get("sections", [])
    body_text = ""
    if sections:
        body_text = sections[0].get("body", "")

    txbox = slide.shapes.add_textbox(Inches(0.8), Inches(1.3), Inches(11.5), Inches(5.5))
    txbox.word_wrap = True
    tf = txbox.text_frame
    tf.word_wrap = True
    _add_rich_body(tf, body_text, font_size=16, color=CLR_DARK)


def _build_finding_slide(prs, finding: dict, evidence_items: list):
    """Slides 3-N — One per finding."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _set_slide_bg(slide)
    _add_purple_bar(slide)

    number = finding.get("number", "?")
    confidence = finding.get("confidence", "medium")

    # Finding number + confidence badge
    badge_text = f"Finding #{number}  —  {confidence.upper()} confidence"
    badge_box = _add_textbox(slide, Inches(0.8), Inches(0.35), Inches(5), Inches(0.45),
                             badge_text, font_size=12, bold=True,
                             color=_confidence_color(confidence))

    # Claim (headline)
    claim = _strip_md(finding.get("claim", ""))
    _add_textbox(slide, Inches(0.8), Inches(0.85), Inches(11.5), Inches(0.9),
                 claim, font_size=22, bold=True, color=CLR_DARK)

    # Support text as bullet points
    support = finding.get("support", "")
    if support:
        # Split on sentence boundaries for bullet points
        sentences = re.split(r"(?<=[.!?])\s+", support)
        sentences = [s.strip() for s in sentences if s.strip()]

        y_pos = Inches(1.9)
        for sentence in sentences:
            bullet_box = slide.shapes.add_textbox(
                Inches(1.0), y_pos, Inches(7.0), Inches(0.4),
            )
            bullet_box.word_wrap = True
            tf = bullet_box.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            run = p.add_run()
            run.text = f"•  {_strip_md(sentence)}"
            _set_font(run, size=13, color=CLR_DARK)
            y_pos += Inches(0.4)

    # Evidence callout box (1-2 quotes)
    if evidence_items:
        callout_top = Inches(4.2)
        callout = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(0.8), callout_top, Inches(11.5), Inches(2.8),
        )
        callout.fill.solid()
        callout.fill.fore_color.rgb = CLR_ACCENT
        callout.line.fill.background()

        _add_textbox(slide, Inches(1.0), callout_top + Inches(0.15),
                     Inches(3), Inches(0.35),
                     "EVIDENCE", font_size=10, bold=True, color=CLR_PRIMARY)

        quote_y = callout_top + Inches(0.5)
        for ev in evidence_items[:2]:
            quote = ev.get("quote", "")
            if len(quote) > 280:
                quote = quote[:277] + "..."
            sentiment = ev.get("sentiment", "")
            label = f'"{quote}"'
            if sentiment:
                label += f"  [{sentiment}]"
            qtb = _add_textbox(slide, Inches(1.2), quote_y,
                               Inches(10.5), Inches(1.0),
                               label, font_size=11, color=CLR_MID_GREY)
            qtb.text_frame.paragraphs[0].runs[0].font.italic = True
            quote_y += Inches(1.1)


def _build_stats_slide(prs, report: dict, tagged_data: list):
    """Key Statistics slide — sentiment table, top themes, top signals."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _set_slide_bg(slide)
    _add_purple_bar(slide)

    _add_textbox(slide, Inches(0.8), Inches(0.4), Inches(11.5), Inches(0.7),
                 "Key Statistics", font_size=28, bold=True, color=CLR_PRIMARY)

    # Gather stats from tagged data
    from collections import Counter
    sentiments = Counter()
    themes = Counter()
    signals = Counter()
    total = 0
    for row in tagged_data:
        if row.get("error"):
            continue
        total += 1
        if row.get("sentiment"):
            sentiments[row["sentiment"]] += 1
        if row.get("theme"):
            themes[row["theme"]] += 1
        sig_val = row.get("signals", "")
        if sig_val and isinstance(sig_val, str):
            for s in sig_val.split(","):
                s = s.strip()
                if s:
                    signals[s] += 1

    # ── Sentiment table ──────────────────────────────────────────────────
    _add_textbox(slide, Inches(0.8), Inches(1.3), Inches(4), Inches(0.4),
                 "Sentiment Distribution", font_size=14, bold=True, color=CLR_DARK)

    sent_items = sentiments.most_common(6)
    if sent_items:
        rows_count = len(sent_items) + 1  # header + data
        table_shape = slide.shapes.add_table(
            rows_count, 3, Inches(0.8), Inches(1.8), Inches(5), Inches(0.35 * rows_count),
        )
        table = table_shape.table

        # Header row
        for ci, hdr in enumerate(["Sentiment", "Count", "%"]):
            cell = table.cell(0, ci)
            cell.text = hdr
            for p in cell.text_frame.paragraphs:
                for r in p.runs:
                    _set_font(r, size=10, bold=True, color=CLR_WHITE)
            cell.fill.solid()
            cell.fill.fore_color.rgb = CLR_PRIMARY

        for ri, (sent, cnt) in enumerate(sent_items, start=1):
            pct = f"{cnt / total * 100:.0f}%" if total else "0%"
            for ci, val in enumerate([sent, str(cnt), pct]):
                cell = table.cell(ri, ci)
                cell.text = val
                for p in cell.text_frame.paragraphs:
                    for r in p.runs:
                        _set_font(r, size=10, color=CLR_DARK)
                cell.fill.solid()
                cell.fill.fore_color.rgb = CLR_WHITE if ri % 2 == 0 else CLR_LIGHT_BG

    # ── Top themes ────────────────────────────────────────────────────────
    _add_textbox(slide, Inches(7.0), Inches(1.3), Inches(5.5), Inches(0.4),
                 "Top Themes", font_size=14, bold=True, color=CLR_DARK)

    theme_y = Inches(1.8)
    for theme, cnt in themes.most_common(7):
        pct = f"{cnt / total * 100:.0f}%" if total else "0%"
        _add_textbox(slide, Inches(7.0), theme_y, Inches(5.5), Inches(0.35),
                     f"•  {theme}  ({cnt}, {pct})", font_size=11, color=CLR_DARK)
        theme_y += Inches(0.35)

    # ── Top signals ───────────────────────────────────────────────────────
    if signals:
        _add_textbox(slide, Inches(7.0), theme_y + Inches(0.3), Inches(5.5), Inches(0.4),
                     "Top Signals", font_size=14, bold=True, color=CLR_DARK)

        sig_y = theme_y + Inches(0.7)
        for sig, cnt in signals.most_common(5):
            _add_textbox(slide, Inches(7.0), sig_y, Inches(5.5), Inches(0.35),
                         f"•  {sig}  ({cnt})", font_size=11, color=CLR_DARK)
            sig_y += Inches(0.35)

    # Total items note
    _add_textbox(slide, Inches(0.8), Inches(6.6), Inches(5), Inches(0.4),
                 f"Total items analyzed: {total}", font_size=11, bold=False,
                 color=CLR_MID_GREY)


def _build_recommendations_slide(prs, report: dict):
    """Recommendations slide — the so_what section."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _set_slide_bg(slide)
    _add_purple_bar(slide)

    _add_textbox(slide, Inches(0.8), Inches(0.4), Inches(11.5), Inches(0.7),
                 "Recommendations & Next Steps", font_size=28, bold=True, color=CLR_PRIMARY)

    so_what = report.get("so_what", "")
    txbox = slide.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(11.5), Inches(5.0))
    txbox.word_wrap = True
    tf = txbox.text_frame
    tf.word_wrap = True
    _add_rich_body(tf, so_what, font_size=16, color=CLR_DARK)

    # Additional sections beyond the first (exec summary) as supplementary content
    sections = report.get("sections", [])
    if len(sections) > 1:
        y_pos = Inches(3.5)
        for section in sections[1:]:
            heading = section.get("heading", "")
            body = section.get("body", "")
            if not heading or not body:
                continue
            _add_textbox(slide, Inches(0.8), y_pos, Inches(11.5), Inches(0.4),
                         heading, font_size=14, bold=True, color=CLR_PRIMARY)
            y_pos += Inches(0.45)
            tb = _add_textbox(slide, Inches(0.8), y_pos, Inches(11.5), Inches(0.6),
                              _strip_md(body), font_size=11, color=CLR_DARK)
            y_pos += Inches(0.7)
            if y_pos > Inches(6.5):
                break


def _build_appendix_slide(prs, tagged_data: list):
    """Appendix — summary table of tagged data (top 20 rows)."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _set_slide_bg(slide)
    _add_purple_bar(slide)

    _add_textbox(slide, Inches(0.8), Inches(0.4), Inches(11.5), Inches(0.7),
                 "Appendix — Tagged Data Sample", font_size=28, bold=True, color=CLR_PRIMARY)

    # Pick columns that are commonly useful
    preferred_cols = [
        "theme", "sentiment", "sentiment_nuance", "emotion",
        "driver", "signals", "severity", "confidence",
        "brand", "primary_brand", "stage",
    ]

    valid_rows = [r for r in tagged_data if not r.get("error")][:20]
    if not valid_rows:
        _add_textbox(slide, Inches(0.8), Inches(1.5), Inches(11.5), Inches(1.0),
                     "No valid tagged data available.", font_size=14, color=CLR_MID_GREY)
        return

    # Determine which columns to show (present in data)
    available_cols = []
    for col in preferred_cols:
        if any(r.get(col) for r in valid_rows):
            available_cols.append(col)
    # If none of the preferred columns match, take first few keys
    if not available_cols:
        sample_keys = list(valid_rows[0].keys())
        skip_keys = {"error", "error_message", "_status"}
        available_cols = [k for k in sample_keys if k not in skip_keys][:6]

    # Limit columns to fit on slide
    available_cols = available_cols[:6]
    num_cols = len(available_cols)
    num_rows = len(valid_rows) + 1  # header + data

    col_width = min(Inches(2.0), (SLIDE_WIDTH - Inches(1.6)) / num_cols)
    table_width = col_width * num_cols

    table_shape = slide.shapes.add_table(
        num_rows, num_cols,
        Inches(0.8), Inches(1.3),
        int(table_width), Inches(min(5.5, 0.3 * num_rows)),
    )
    table = table_shape.table

    # Header
    for ci, col_name in enumerate(available_cols):
        cell = table.cell(0, ci)
        cell.text = col_name.replace("_", " ").title()
        for p in cell.text_frame.paragraphs:
            for r in p.runs:
                _set_font(r, size=8, bold=True, color=CLR_WHITE)
        cell.fill.solid()
        cell.fill.fore_color.rgb = CLR_PRIMARY

    # Data rows
    for ri, row in enumerate(valid_rows, start=1):
        for ci, col_name in enumerate(available_cols):
            cell = table.cell(ri, ci)
            val = row.get(col_name, "")
            cell_text = str(val) if val is not None else ""
            if len(cell_text) > 50:
                cell_text = cell_text[:47] + "..."
            cell.text = cell_text
            for p in cell.text_frame.paragraphs:
                for r in p.runs:
                    _set_font(r, size=7, color=CLR_DARK)
            cell.fill.solid()
            cell.fill.fore_color.rgb = CLR_WHITE if ri % 2 == 0 else CLR_LIGHT_BG

    note = f"Showing {len(valid_rows)} of {len([r for r in tagged_data if not r.get('error')])} items"
    _add_textbox(slide, Inches(0.8), Inches(6.9), Inches(6), Inches(0.35),
                 note, font_size=9, color=CLR_MID_GREY)


# ── Public API ────────────────────────────────────────────────────────────────

def build_report_pptx(report_data: dict, tagged_data: list) -> bytes:
    """Generate a polished PPTX presentation from report JSON and tagged data.

    Parameters
    ----------
    report_data : dict
        Structured report with title, subtitle, sections, findings, evidence,
        so_what fields (as returned by the generate-report endpoint).
    tagged_data : list
        Raw tagged / analyzed data rows.

    Returns
    -------
    bytes
        The PPTX file contents.
    """
    prs = Presentation()
    prs.slide_width  = SLIDE_WIDTH
    prs.slide_height = SLIDE_HEIGHT

    # 1. Title slide
    _build_title_slide(prs, report_data)

    # 2. Executive Summary
    _build_exec_summary_slide(prs, report_data)

    # 3-N. Finding slides
    findings = report_data.get("findings", [])
    evidence_list = report_data.get("evidence", [])
    for i, finding in enumerate(findings):
        # Assign 1-2 evidence items to each finding
        ev_start = i * 2
        ev_slice = evidence_list[ev_start:ev_start + 2]
        _build_finding_slide(prs, finding, ev_slice)

    # N+1. Key Statistics
    _build_stats_slide(prs, report_data, tagged_data)

    # N+2. Recommendations
    _build_recommendations_slide(prs, report_data)

    # Last. Appendix
    _build_appendix_slide(prs, tagged_data)

    buf = io.BytesIO()
    prs.save(buf)
    buf.seek(0)
    return buf.getvalue()


def build_data_pptx(tagged_data: list, filename: str = "Data") -> bytes:
    """Generate a simpler PPTX that presents tagged data as slide tables.

    Splits data across multiple slides if there are many rows (max ~15 rows
    per slide for readability).

    Parameters
    ----------
    tagged_data : list
        The analyzed / tagged data rows.
    filename : str
        Original file name for the title slide.

    Returns
    -------
    bytes
        The PPTX file contents.
    """
    ROWS_PER_SLIDE = 15

    prs = Presentation()
    prs.slide_width  = SLIDE_WIDTH
    prs.slide_height = SLIDE_HEIGHT

    # Title slide
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _set_slide_bg(slide, CLR_DARK)
    block = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0), Inches(0),
        SLIDE_WIDTH, Inches(3.0),
    )
    block.fill.solid()
    block.fill.fore_color.rgb = CLR_PRIMARY
    block.line.fill.background()

    _add_textbox(slide, Inches(0.8), Inches(0.8), Inches(11.5), Inches(1.2),
                 f"{filename} — Tagged Data Export",
                 font_size=32, bold=True, color=CLR_WHITE)
    _add_textbox(slide, Inches(0.8), Inches(2.1), Inches(11.5), Inches(0.6),
                 f"{len(tagged_data)} rows  |  Generated {datetime.utcnow().strftime('%d %B %Y')}",
                 font_size=16, color=CLR_WHITE)

    # Determine columns
    valid_rows = [r for r in tagged_data if not r.get("error")]
    if not valid_rows:
        slide2 = prs.slides.add_slide(prs.slide_layouts[6])
        _set_slide_bg(slide2)
        _add_textbox(slide2, Inches(2), Inches(3), Inches(9), Inches(1),
                     "No valid data to display.", font_size=20, color=CLR_MID_GREY,
                     alignment=PP_ALIGN.CENTER)
        buf = io.BytesIO()
        prs.save(buf)
        buf.seek(0)
        return buf.getvalue()

    # Choose columns: skip internal keys, limit to 8 for readability
    skip_keys = {"error", "error_message", "_status"}
    all_keys = list(valid_rows[0].keys())
    display_cols = [k for k in all_keys if k not in skip_keys][:8]
    num_cols = len(display_cols)

    # Chunk rows into slides
    for chunk_start in range(0, len(valid_rows), ROWS_PER_SLIDE):
        chunk = valid_rows[chunk_start:chunk_start + ROWS_PER_SLIDE]
        num_data_rows = len(chunk)
        total_rows = num_data_rows + 1  # header

        slide = prs.slides.add_slide(prs.slide_layouts[6])
        _set_slide_bg(slide)
        _add_purple_bar(slide)

        page_num = (chunk_start // ROWS_PER_SLIDE) + 1
        total_pages = (len(valid_rows) + ROWS_PER_SLIDE - 1) // ROWS_PER_SLIDE
        _add_textbox(slide, Inches(0.8), Inches(0.3), Inches(11.5), Inches(0.5),
                     f"Data — Page {page_num} of {total_pages}",
                     font_size=18, bold=True, color=CLR_PRIMARY)

        col_width = min(Inches(2.0), (SLIDE_WIDTH - Inches(1.6)) / num_cols)
        table_width = col_width * num_cols
        row_height = min(Inches(0.35), Inches(6.0) / total_rows)

        table_shape = slide.shapes.add_table(
            total_rows, num_cols,
            Inches(0.8), Inches(0.9),
            int(table_width), int(row_height * total_rows),
        )
        table = table_shape.table

        # Header
        for ci, col_name in enumerate(display_cols):
            cell = table.cell(0, ci)
            cell.text = col_name.replace("_", " ").title()
            for p in cell.text_frame.paragraphs:
                for r in p.runs:
                    _set_font(r, size=8, bold=True, color=CLR_WHITE)
            cell.fill.solid()
            cell.fill.fore_color.rgb = CLR_PRIMARY

        # Data
        for ri, row in enumerate(chunk, start=1):
            for ci, col_name in enumerate(display_cols):
                cell = table.cell(ri, ci)
                val = row.get(col_name, "")
                cell_text = str(val) if val is not None else ""
                if len(cell_text) > 60:
                    cell_text = cell_text[:57] + "..."
                cell.text = cell_text
                for p in cell.text_frame.paragraphs:
                    for r in p.runs:
                        _set_font(r, size=7, color=CLR_DARK)
                cell.fill.solid()
                cell.fill.fore_color.rgb = CLR_WHITE if ri % 2 == 0 else CLR_LIGHT_BG

    buf = io.BytesIO()
    prs.save(buf)
    buf.seek(0)
    return buf.getvalue()
